from openpyxl import Workbook
from openpyxl.styles import Font


def generate_attendance_excel(records, filename):

    wb = Workbook()

    ws = wb.active
    ws.title = "Attendance"

    headers = [
        "Date",
        "Attendance Type",
        "Event",
        "Member Name",
        "ITS Number",
        "Status",
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)

    row = 2

    for record in records:

        ws.cell(row=row, column=1, value=record["date"])
        ws.cell(row=row, column=2, value=record["attendance_type"])
        ws.cell(row=row, column=3, value=record.get("event_name"))
        ws.cell(row=row, column=4, value=record["name"])
        ws.cell(row=row, column=5, value=record["its_no"])
        ws.cell(row=row, column=6, value=record["status"])

        row += 1

    wb.save(filename)

    return filename
